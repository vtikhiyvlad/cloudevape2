"""
Умная загрузка товаров из Excel.
Ожидаемые колонки:
Категория | Название | Крепость | Вкус | Цена | Остаток
"""
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.models.product import Category, Product, ProductVariant
from app.models.inventory import StockMovement
from decimal import Decimal
from io import BytesIO


EXPECTED_COLUMNS = {
    "категория": "category",
    "название": "name",
    "крепость": "strength",
    "вкус": "flavor",
    "цена": "price",
    "остаток": "stock",
}


def normalize_header(value: str) -> str | None:
    if not value:
        return None
    cleaned = str(value).strip().lower()
    return EXPECTED_COLUMNS.get(cleaned)


async def import_products_from_excel(
    session: AsyncSession,
    store_id: int,
    user_id: int,
    file_bytes: bytes,
) -> dict:
    """
    Импортирует/обновляет товары.
    Возвращает статистику: создано, обновлено, ошибки.
    """
    wb = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"error": "Файл пустой"}

    # Определяем заголовки
    header_row = rows[0]
    col_map = {}
    for idx, cell in enumerate(header_row):
        key = normalize_header(cell)
        if key:
            col_map[key] = idx

    required = ["category", "name", "price"]
    missing = [r for r in required if r not in col_map]
    if missing:
        return {"error": f"Не хватает колонок: {', '.join(missing)}. Нужны: Категория, Название, Цена, (Крепость, Вкус, Остаток)"}

    stats = {
        "created_products": 0,
        "created_variants": 0,
        "updated_variants": 0,
        "errors": [],
        "total_rows": 0,
    }

    # Кэш категорий
    cat_result = await session.execute(
        select(Category).where(Category.store_id == store_id)
    )
    categories = {c.name.lower(): c for c in cat_result.scalars().all()}

    for row_idx, row in enumerate(rows[1:], start=2):
        try:
            def get(col):
                idx = col_map.get(col)
                if idx is None or idx >= len(row):
                    return None
                val = row[idx]
                return str(val).strip() if val is not None else None

            category_name = get("category")
            product_name = get("name")
            strength = get("strength") or None
            flavor = get("flavor") or None
            price_raw = get("price")
            stock_raw = get("stock")

            if not category_name or not product_name:
                continue

            stats["total_rows"] += 1

            # Цена
            try:
                price = Decimal(str(price_raw).replace(",", ".").replace(" ", ""))
            except Exception:
                stats["errors"].append(f"Строка {row_idx}: неверная цена '{price_raw}'")
                continue

            # Остаток
            stock = 0
            if stock_raw is not None and str(stock_raw).strip() != "":
                try:
                    stock = int(float(str(stock_raw).replace(",", ".")))
                except Exception:
                    stats["errors"].append(f"Строка {row_idx}: неверный остаток '{stock_raw}'")
                    continue

            # Категория
            cat_key = category_name.lower()
            if cat_key not in categories:
                cat = Category(store_id=store_id, name=category_name)
                session.add(cat)
                await session.flush()
                categories[cat_key] = cat
            category = categories[cat_key]

            # Товар
            prod_result = await session.execute(
                select(Product).where(
                    Product.store_id == store_id,
                    Product.name == product_name,
                    Product.category_id == category.id,
                )
            )
            product = prod_result.scalar_one_or_none()
            if not product:
                product = Product(
                    store_id=store_id,
                    category_id=category.id,
                    name=product_name,
                )
                session.add(product)
                await session.flush()
                stats["created_products"] += 1

            # Вариант
            var_result = await session.execute(
                select(ProductVariant).where(
                    ProductVariant.product_id == product.id,
                    ProductVariant.strength == strength,
                    ProductVariant.flavor == flavor,
                )
            )
            variant = var_result.scalar_one_or_none()

            if variant:
                old_stock = variant.stock
                variant.price = price
                if stock_raw is not None:
                    # Обновляем остаток только если явно указан
                    variant.stock = stock
                    if stock != old_stock:
                        movement = StockMovement(
                            variant_id=variant.id,
                            user_id=user_id,
                            store_id=store_id,
                            movement_type="import",
                            quantity_change=stock - old_stock,
                            stock_before=old_stock,
                            stock_after=stock,
                            comment="Обновление через Excel",
                        )
                        session.add(movement)
                stats["updated_variants"] += 1
            else:
                variant = ProductVariant(
                    product_id=product.id,
                    strength=strength,
                    flavor=flavor,
                    price=price,
                    stock=stock,
                )
                session.add(variant)
                await session.flush()

                if stock > 0:
                    movement = StockMovement(
                        variant_id=variant.id,
                        user_id=user_id,
                        store_id=store_id,
                        movement_type="initial",
                        quantity_change=stock,
                        stock_before=0,
                        stock_after=stock,
                        comment="Первая загрузка через Excel",
                    )
                    session.add(movement)
                stats["created_variants"] += 1

        except Exception as e:
            stats["errors"].append(f"Строка {row_idx}: {str(e)}")

    await session.flush()
    return stats
